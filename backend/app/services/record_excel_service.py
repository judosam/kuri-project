"""Admin full-record Excel export, including inspection answers, Old Unique ID, and photos."""
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.google_drive_service import GoogleDriveService
from app.services.local_storage_service import LocalStorageService

REPORT_BATCH_SIZE = 500

# ─── Modern Spreadsheet Styling Theme ─────────────────────────────────────────
FONT_FAMILY = "Segoe UI"
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E293B")  # Slate 800 executive navy

ROW_FILL_ODD = PatternFill("solid", fgColor="FFFFFF")
ROW_FILL_EVEN = PatternFill("solid", fgColor="F8FAFC")  # Subtle Slate 50 tint

THIN_BORDER_SIDE = Side(border_style="thin", color="E2E8F0")
DATA_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)

HEADER_BORDER_BOTTOM = Side(border_style="medium", color="0F172A")
HEADER_BORDER_SIDE = Side(border_style="thin", color="334155")
HEADER_BORDER = Border(
    left=HEADER_BORDER_SIDE,
    right=HEADER_BORDER_SIDE,
    top=HEADER_BORDER_SIDE,
    bottom=HEADER_BORDER_BOTTOM,
)

DATA_FONT = Font(name=FONT_FAMILY, size=10, color="0F172A")
DATA_FONT_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
DATA_FONT_MUTED = Font(name=FONT_FAMILY, size=10, color="64748B")

ALIGN_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=False)
ALIGN_LEFT = Alignment(vertical="center", horizontal="left", wrap_text=True)
ALIGN_RIGHT = Alignment(vertical="center", horizontal="right", wrap_text=False)
ALIGN_HEADER = Alignment(vertical="center", horizontal="center", wrap_text=True)

# Fields that should be centered in the modern layout
CENTER_ALIGNED_FIELDS = {
    "id", "unique_id", "old_unique_id", "source", "priority_tier",
    "govt_dtcp_permit", "dtcp_permit_date", "dtcp_application_no",
    "survey_no", "subdivision_no", "status", "qc_status", "audit_status",
    "telecaller_status", "revisit_status", "date_of_visit",
    "assigned_at", "assigned_date", "field_captured_at", "captured_at",
    "audited_at", "qc_completed_at", "qc_edited_at", "telecaller_called_at",
    "reassigned_at", "updated_at", "last_modified_at", "created_at",
    "follow_up_date", "no_of_floors", "built_up_area", "qc_sqfeet",
    "photo_count", "field_latitude", "field_longitude",
    "field_location_accuracy", "deviation_meters", "coordinates",
    "updated_coordinates", "telecaller_no_of_floors", "telecaller_built_up_area",
    "construction_happening", "contact_number", "site_contractor_architect_number",
    "qc_builder_contact",
}

# Optimal column widths for a clean modern spreadsheet
COLUMN_WIDTH_MAP = {
    "id": 13,
    "unique_id": 16,
    "old_unique_id": 24,
    "source": 18,
    "last_modified_at": 20,
    "name": 25,
    "created_at": 20,
    "coordinates": 24,
    "area": 18,
    "village": 22,
    "taluk": 20,
    "district": 16,
    "site_address": 38,
    "street_address": 38,
    "admin_formatted_address": 38,
    "priority_tier": 15,
    "govt_dtcp_permit": 18,
    "dtcp_application_no": 22,
    "dtcp_permit_date": 18,
    "survey_no": 15,
    "subdivision_no": 15,
    "land_use": 24,
    "owner_name_en": 25,
    "staged_operator_name": 22,
    "uploaded_by_name": 20,
    "assigned_auditor_name": 18,
    "qc_status": 18,
    "audited_by_name": 18,
    "audited_at": 20,
    "audit_comment": 36,
    "qc_vhr_comment": 36,
    "qc_street_details": 36,
    "qc_builder_contact": 22,
    "qc_sqfeet": 14,
    "qc_completed_at": 20,
    "qc_edited_at": 20,
    "qc_by_name": 18,
    "user_name": 20,
    "operator_email": 25,
    "status": 20,
    "assigned_at": 20,
    "date_of_visit": 18,
    "field_captured_at": 22,
    "formatted_address": 38,
    "field_latitude": 16,
    "field_longitude": 16,
    "field_location_accuracy": 15,
    "deviation_meters": 18,
    "construction_happening": 22,
    "type_of_construction": 24,
    "stage_of_construction": 24,
    "no_of_floors": 14,
    "built_up_area": 16,
    "expected_completion": 20,
    "relationship_with_building": 24,
    "building_user_name": 24,
    "spoken_with": 20,
    "contact_number": 18,
    "site_contractor_architect_name": 25,
    "site_contractor_architect_number": 18,
    "remarks": 40,
    "photo_count": 14,
    "captured_at": 24,
    "assigned_telecaller_name": 20,
    "telecaller_status": 18,
    "telecaller_comment": 36,
    "telecaller_called_at": 20,
    "telecaller_by_name": 18,
    "telecaller_stage_of_construction": 24,
    "telecaller_no_of_floors": 16,
    "telecaller_built_up_area": 16,
    "telecaller_expected_completion": 22,
    "reassigned_to_name": 22,
    "reassigned_at": 20,
    "revisit_status": 18,
    "updated_coordinates": 24,
    "updated_information": 36,
    "follow_up_date": 18,
    "updated_by_role": 18,
    "updated_by_name": 20,
    "updated_at": 20,
}

# ─── Old Unique ID Reference Cache ───────────────────────────────────────────
_OLD_UNIQUE_ID_CACHE = None


def _get_old_unique_id_map() -> dict[str, str]:
    """Load the complete mapping from Updated_Unique_IDs_Reference.xlsx into memory."""
    global _OLD_UNIQUE_ID_CACHE
    if _OLD_UNIQUE_ID_CACHE is not None:
        return _OLD_UNIQUE_ID_CACHE

    mapping: dict[str, str] = {}
    candidate_paths = [
        Path(__file__).parent.parent.parent / "data" / "Updated_Unique_IDs_Reference.xlsx",
        Path("backend/data/Updated_Unique_IDs_Reference.xlsx"),
        Path("data/Updated_Unique_IDs_Reference.xlsx"),
        Path("/opt/photocapture/backend/data/Updated_Unique_IDs_Reference.xlsx"),
    ]
    for p in candidate_paths:
        if p.exists():
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(p), data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True, min_row=2):
                    rec_id, old_uid, new_uid = row[0], row[1], row[2]
                    old_str = str(old_uid).strip() if old_uid else ""
                    new_str = str(new_uid).strip() if new_uid else ""
                    if new_str and old_str:
                        mapping[new_str] = old_str
                    if rec_id is not None and old_str:
                        mapping[str(rec_id)] = old_str
                break
            except Exception:
                pass
    _OLD_UNIQUE_ID_CACHE = mapping
    return _OLD_UNIQUE_ID_CACHE


def iter_report_records(service, **filters):
    """Read every matching record, rather than silently stopping at 10,000."""
    page = 1
    while True:
        records, total = service.get_all_records(
            page=page, page_size=REPORT_BATCH_SIZE, sort_by='id', sort_order='asc', **filters
        )
        yield from records
        if not records or page * REPORT_BATCH_SIZE >= total:
            return
        page += 1


def _photo_bytes(record, photo):
    key = photo.google_drive_file_id
    if key.startswith(('pending:', 'simulated-')):
        return None
    storage = LocalStorageService()
    path = storage.get_file_path(key) or storage.get_file_path(
        f'{record.user_id}/{record.id}/{photo.file_name}'
    )
    if path:
        return Path(path).read_bytes()
    if not any(separator in key for separator in ('/', '\\', ':')):
        drive = GoogleDriveService()
        if drive.service:
            return drive.service.files().get_media(fileId=key).execute()
    return None


def _format_cell_value(field: str, val):
    """Normalize and format cell values for clean, executive-ready presentation."""
    if val is None or val == "" or val == "__EMPTY__" or val == "__BLANK__" or val == "(Blanks)":
        return "—"

    # Booleans
    if isinstance(val, bool):
        return "Yes" if val else "No"

    # Keep source matching stored value for system and test compatibility
    if field == "source":
        return val

    # Captured at (UTC) format for unit test compatibility
    if field == "captured_at":
        if isinstance(val, datetime):
            return val.strftime("%Y-%m-%dT%H:%M:%SZ")
        val_str = str(val).strip()
        if val_str and not val_str.endswith("Z"):
            return val_str + "Z"
        return val_str

    # Date formatting for other timestamps
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M")

    if field in ("created_at", "last_modified_at", "assigned_at", "field_captured_at",
                 "audited_at", "qc_completed_at", "qc_edited_at", "telecaller_called_at",
                 "reassigned_at", "updated_at"):
        val_str = str(val).strip()
        if "T" in val_str or (len(val_str) >= 19 and "-" in val_str and ":" in val_str):
            try:
                dt = datetime.fromisoformat(val_str.replace("Z", "+00:00"))
                return dt.strftime("%Y-%m-%d %H:%M")
            except Exception:
                pass
        return val_str

    return val


def build_record_excel(record, detail: dict) -> bytes:
    """Modernized single-record Excel export."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Record Details'

    # Ensure Old Unique ID is populated
    old_uid = (
        detail.get("old_unique_id")
        or getattr(record, "old_unique_id", None)
        or _get_old_unique_id_map().get(str(detail.get("unique_id") or ""))
        or _get_old_unique_id_map().get(str(record.id))
        or "—"
    )
    detail["old_unique_id"] = old_uid

    labels = {
        'id': 'Record ID',
        'unique_id': 'Unique ID',
        'old_unique_id': 'Old Unique ID',
        'user_id': 'User ID',
        'owner_name_en': 'Owner / Client Name',
        'user_name': 'Field Verifier',
        'assigned_auditor_name': 'QA User',
        'created_at': 'Created At',
        'source': 'Record Source',
        'last_modified_at': 'Last Modified Date',
    }

    sheet.append(['Field', 'Value'])
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = ALIGN_HEADER

    row_count = 1
    for field, value in detail.items():
        if field in ('images', 'errors'):
            continue
        field_label = labels.get(field, field.replace('_', ' ').title())
        formatted_val = _format_cell_value(field, value)
        sheet.append([field_label, formatted_val])
        row_count += 1
        r_fill = ROW_FILL_EVEN if (row_count % 2 == 0) else ROW_FILL_ODD
        cell_a = sheet.cell(row_count, 1)
        cell_b = sheet.cell(row_count, 2)
        cell_a.font = DATA_FONT_BOLD
        cell_a.fill = r_fill
        cell_a.border = DATA_BORDER
        cell_a.alignment = ALIGN_LEFT
        cell_b.font = DATA_FONT
        cell_b.fill = r_fill
        cell_b.border = DATA_BORDER
        cell_b.alignment = ALIGN_LEFT
        if isinstance(cell_b.value, str):
            cell_b.data_type = 's'
        val_len = len(str(formatted_val or ''))
        sheet.row_dimensions[row_count].height = max(26, min(140, 20 * (1 + val_len // 80)))

    # Extra summary rows
    operator_email = record.user.email if record.user else "—"
    photo_count = len(record.images)
    for extra_field, extra_val in [('Operator Email', operator_email), ('Photo Count', photo_count)]:
        sheet.append([extra_field, extra_val])
        row_count += 1
        r_fill = ROW_FILL_EVEN if (row_count % 2 == 0) else ROW_FILL_ODD
        cell_a = sheet.cell(row_count, 1)
        cell_b = sheet.cell(row_count, 2)
        cell_a.font = DATA_FONT_BOLD
        cell_a.fill = r_fill
        cell_a.border = DATA_BORDER
        cell_a.alignment = ALIGN_LEFT
        cell_b.font = DATA_FONT
        cell_b.fill = r_fill
        cell_b.border = DATA_BORDER
        cell_b.alignment = ALIGN_LEFT
        sheet.row_dimensions[row_count].height = 26

    sheet.column_dimensions['A'].width = 34
    sheet.column_dimensions['B'].width = 95
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = True

    # Photos sheet
    photos = workbook.create_sheet('Photos')
    photos.append([
        'Photo', 'File Name', 'Source', 'Captured At (UTC)',
        'Latitude', 'Longitude', 'Accuracy (m)', 'Image / Availability'
    ])
    photos.row_dimensions[1].height = 34
    for cell in photos[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = ALIGN_HEADER

    if not record.images:
        photos.append(['No photos attached to this record.'])
        p_row = photos.max_row
        r_fill = ROW_FILL_ODD
        for c in photos[p_row]:
            c.font = DATA_FONT_MUTED
            c.fill = r_fill
            c.border = DATA_BORDER
        photos.row_dimensions[p_row].height = 26
    else:
        for photo in record.images:
            photos.append([
                photo.photo_order, photo.file_name, photo.source,
                photo.captured_at.isoformat() if photo.captured_at else "—",
                photo.latitude, photo.longitude, photo.location_accuracy, 'Image unavailable'
            ])
            row = photos.max_row
            r_fill = ROW_FILL_EVEN if (row % 2 == 0) else ROW_FILL_ODD
            for c in photos[row]:
                c.font = DATA_FONT
                c.fill = r_fill
                c.border = DATA_BORDER
                c.alignment = ALIGN_CENTER
                if isinstance(c.value, str):
                    c.data_type = 's'
            photos.row_dimensions[row].height = 28
            try:
                content = _photo_bytes(record, photo)
                if content:
                    image = ExcelImage(io.BytesIO(content))
                    scale = min(1, 480 / image.width, 300 / image.height)
                    image.width *= scale
                    image.height *= scale
                    photos.add_image(image, f'H{row}')
                    photos.cell(row, 8).value = None
                    photos.row_dimensions[row].height = max(40, image.height * 0.75 + 12)
            except Exception:
                photos.cell(row, 8).value = 'Image unavailable'

    for col in ('A', 'E', 'F', 'G'):
        photos.column_dimensions[col].width = 18
    for col in ('B', 'C', 'D'):
        photos.column_dimensions[col].width = 28
    photos.column_dimensions['H'].width = 70
    photos.freeze_panes = 'A2'
    photos.sheet_view.showGridLines = True
    photos.auto_filter.ref = photos.dimensions

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# ─── Normalized Column Grouping ───────────────────────────────────────────────
SECTION_FIELD_MAP = {
    # 1. Imported Data (Site Planning & Core Identifiers)
    "planning": [
        "coordinates", "area", "village", "taluk", "district",
        "site_address", "street_address", "admin_formatted_address",
        "priority_tier", "govt_dtcp_permit", "dtcp_application_no", "dtcp_permit_date",
        "survey_no", "subdivision_no", "land_use", "owner_name_en",
        "staged_operator_name", "uploaded_by_name",
    ],
    # 2. DC / QA Review
    "audit": [
        "assigned_auditor_name", "qc_status", "audited_by_name", "audited_at",
        "audit_comment", "qc_vhr_comment", "qc_street_details",
        "qc_builder_contact", "qc_sqfeet", "qc_completed_at", "qc_edited_at", "qc_by_name",
    ],
    # 3. Field Verifier Data
    "field": [
        "user_name", "operator_email", "status", "assigned_at",
        "date_of_visit", "field_captured_at", "formatted_address",
        "field_latitude", "field_longitude", "field_location_accuracy",
        "deviation_meters",
        "construction_happening", "type_of_construction", "stage_of_construction",
        "no_of_floors", "built_up_area", "expected_completion",
        "relationship_with_building", "building_user_name", "spoken_with", "contact_number",
        "site_contractor_architect_name", "site_contractor_architect_number",
        "remarks", "photo_count", "captured_at",
    ],
    # 4. Telecaller Lead Management
    "telecaller": [
        "assigned_telecaller_name", "telecaller_status", "telecaller_comment",
        "telecaller_called_at", "telecaller_by_name",
        "telecaller_stage_of_construction", "telecaller_no_of_floors",
        "telecaller_built_up_area", "telecaller_expected_completion",
    ],
    # 5. Reassign & Revisit
    "reassign": [
        "reassigned_to_name", "reassigned_at", "revisit_status",
        "updated_coordinates", "updated_information", "follow_up_date",
        "updated_by_role", "updated_by_name", "updated_at",
    ],
}

REPORT_LABELS = {
    # Core Identifiers
    "id": "Record ID",
    "unique_id": "Unique ID",
    "old_unique_id": "Old Unique ID",
    "source": "Record Source",
    "last_modified_at": "Last Modified Date",
    "name": "Record Name",
    "created_at": "Created Date",
    # Planning
    "coordinates": "Imported Coordinates",
    "area": "Area",
    "village": "Village",
    "taluk": "Taluk",
    "district": "District",
    "site_address": "Site Address",
    "street_address": "Street Address",
    "admin_formatted_address": "Admin Resolved Address",
    "priority_tier": "Priority Tier",
    "govt_dtcp_permit": "Govt DTCP Permit",
    "dtcp_application_no": "DTCP Application No",
    "dtcp_permit_date": "DTCP Permit Date",
    "survey_no": "Survey No",
    "subdivision_no": "Subdivision No",
    "land_use": "Land Use",
    "owner_name_en": "Owner / Client Name",
    "staged_operator_name": "Suggested Field Verifier",
    "uploaded_by_name": "Imported By",
    # QA / Audit
    "assigned_auditor_name": "QA User",
    "qc_status": "QC Status",
    "audited_by_name": "Audited By",
    "audited_at": "Audited At",
    "audit_comment": "Audit Comments",
    "qc_vhr_comment": "QC Comments (VHR Imagery)",
    "qc_street_details": "QC Additional Details (Street View)",
    "qc_builder_contact": "QC Builder / Contact",
    "qc_sqfeet": "Sq.ft",
    "qc_completed_at": "QC Completed At",
    "qc_edited_at": "QC Edited At",
    "qc_by_name": "QC Done By",
    # Field Verification
    "user_name": "Field Verifier",
    "operator_email": "Operator Email",
    "status": "Verification Status",
    "assigned_at": "Assigned Date",
    "date_of_visit": "Date of Visit",
    "field_captured_at": "Inspection Completed Date",
    "formatted_address": "Field Inspection Address",
    "field_latitude": "Field Latitude",
    "field_longitude": "Field Longitude",
    "field_location_accuracy": "Accuracy (m)",
    "deviation_meters": "GPS Deviation (m)",
    "construction_happening": "Construction Happening",
    "type_of_construction": "Type of Construction",
    "stage_of_construction": "Stage of Construction",
    "no_of_floors": "No. of Floors",
    "built_up_area": "Built-up Area",
    "expected_completion": "Expected Completion",
    "relationship_with_building": "Relationship with Building",
    "building_user_name": "Building User / Resident",
    "spoken_with": "Spoken With",
    "contact_number": "Contact Number",
    "site_contractor_architect_name": "Contractor / Architect",
    "site_contractor_architect_number": "Contractor Contact",
    "remarks": "Inspection Remarks",
    "photo_count": "Photo Count",
    "captured_at": "Captured At (UTC)",
    # Telecaller
    "assigned_telecaller_name": "Assigned Telecaller",
    "telecaller_status": "Telecaller Status",
    "telecaller_comment": "Telecaller Comment",
    "telecaller_called_at": "Telecaller Called Date",
    "telecaller_by_name": "Telecaller Agent",
    "telecaller_stage_of_construction": "Telecaller - Stage of Construction",
    "telecaller_no_of_floors": "Telecaller - No. of Floors",
    "telecaller_built_up_area": "Telecaller - Built-up Area",
    "telecaller_expected_completion": "Telecaller - Expected Completion",
    # Reassign
    "reassigned_to_name": "Reassigned Field Verifier",
    "reassigned_at": "Reassign Date",
    "revisit_status": "Revisit Status",
    "updated_coordinates": "Updated Coordinates",
    "updated_information": "Updated Information",
    "follow_up_date": "Follow-up Date",
    "updated_by_role": "Data Updated By (Role)",
    "updated_by_name": "Data Updated By (Name)",
    "updated_at": "Last Updated Date",
}


def build_records_report(
    records,
    serialize,
    include_images: bool = True,
    sections: list[str] | None = None,
) -> bytes:
    """One row per record with normalized sections, Old Unique ID, and modern styling."""
    # 1. Base Core Identification Columns
    selected_fields: list[str] = [
        "id",
        "unique_id",
        "old_unique_id",
        "source",
        "last_modified_at",
        "name",
        "created_at",
    ]

    # 2. Append fields for active sections
    active_sections = (
        sections
        if sections and any(s in SECTION_FIELD_MAP for s in sections)
        else ["planning", "audit", "field", "telecaller", "reassign"]
    )

    for sec in active_sections:
        if sec in SECTION_FIELD_MAP:
            for f in SECTION_FIELD_MAP[sec]:
                if f not in selected_fields:
                    selected_fields.append(f)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inspections Report"

    # Header Row
    header_labels = [REPORT_LABELS.get(f, f.replace("_", " ").title()) for f in selected_fields]
    sheet.append(header_labels)
    sheet.row_dimensions[1].height = 36

    for col_idx, cell in enumerate(sheet[1], start=1):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = ALIGN_HEADER
        col_letter = get_column_letter(col_idx)
        field_key = selected_fields[col_idx - 1]
        sheet.column_dimensions[col_letter].width = COLUMN_WIDTH_MAP.get(field_key, 22)

    photos_sheet = None
    if include_images:
        photos_sheet = workbook.create_sheet("Photos")
        photos_header = [
            "Record ID", "Unique ID", "Photo ID", "Photo Order", "File Name", "Source",
            "Captured At (UTC)", "Added At (UTC)", "Latitude", "Longitude",
            "Accuracy (m)", "Image / Availability",
        ]
        photos_sheet.append(photos_header)
        photos_sheet.row_dimensions[1].height = 36
        for col_idx, cell in enumerate(photos_sheet[1], start=1):
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = HEADER_BORDER
            cell.alignment = ALIGN_HEADER
            col_letter = get_column_letter(col_idx)
            photos_sheet.column_dimensions[col_letter].width = 20
        photos_sheet.column_dimensions["L"].width = 70

    old_uid_map = _get_old_unique_id_map()

    for record in records:
        detail = serialize(record).model_dump(mode="json")
        images = detail.get("images", [])
        detail["operator_email"] = record.user.email if record.user else None
        detail["photo_count"] = len(images)
        detail["captured_at"] = images[0]["captured_at"] if images else None

        # Resolve Old Unique ID: from DB column, then reference map fallback
        old_uid = (
            detail.get("old_unique_id")
            or getattr(record, "old_unique_id", None)
            or old_uid_map.get(str(detail.get("unique_id") or "").strip())
            or old_uid_map.get(str(record.id))
            or "—"
        )
        detail["old_unique_id"] = old_uid

        if "reassigned_to_name" not in detail or not detail["reassigned_to_name"]:
            detail["reassigned_to_name"] = record.reassigned_to.full_name if getattr(record, "reassigned_to", None) else None
        if "reassigned_at" not in detail or not detail["reassigned_at"]:
            detail["reassigned_at"] = record.reassigned_at.isoformat() if getattr(record, "reassigned_at", None) else None
        if "revisit_status" not in detail or not detail["revisit_status"]:
            detail["revisit_status"] = getattr(record, "revisit_status", None)
        if "updated_coordinates" not in detail or not detail["updated_coordinates"]:
            detail["updated_coordinates"] = getattr(record, "updated_coordinates", None)
        if "updated_information" not in detail or not detail["updated_information"]:
            detail["updated_information"] = getattr(record, "updated_information", None)
        if "updated_by_role" not in detail or not detail["updated_by_role"]:
            detail["updated_by_role"] = getattr(record, "updated_by_role", None)
        if "updated_by_name" not in detail or not detail["updated_by_name"]:
            detail["updated_by_name"] = getattr(record, "updated_by_name", None)
        if "updated_at" not in detail or not detail["updated_at"]:
            up_at = getattr(record, "updated_at", None)
            detail["updated_at"] = up_at.isoformat() if hasattr(up_at, "isoformat") else (str(up_at) if up_at else None)
        if "no_of_floors" not in detail or not detail["no_of_floors"]:
            detail["no_of_floors"] = getattr(record, "no_of_floors", None)
        if "built_up_area" not in detail or not detail["built_up_area"]:
            detail["built_up_area"] = getattr(record, "built_up_area", None)
        if "expected_completion" not in detail or not detail["expected_completion"]:
            detail["expected_completion"] = getattr(record, "expected_completion", None)
        detail["telecaller_no_of_floors"] = getattr(record, "telecaller_no_of_floors", None)
        detail["telecaller_built_up_area"] = getattr(record, "telecaller_built_up_area", None)
        detail["telecaller_expected_completion"] = getattr(record, "telecaller_expected_completion", None)
        detail["telecaller_stage_of_construction"] = getattr(record, "telecaller_stage_of_construction", None)

        # Telecaller data fallbacks
        if "assigned_telecaller_name" not in detail or not detail["assigned_telecaller_name"]:
            detail["assigned_telecaller_name"] = (
                (record.assigned_telecaller.full_name or record.assigned_telecaller.email)
                if getattr(record, "assigned_telecaller", None)
                else None
            )
        if "telecaller_status" not in detail or not detail["telecaller_status"]:
            detail["telecaller_status"] = getattr(record, "telecaller_status", None)
        if "telecaller_comment" not in detail or not detail["telecaller_comment"]:
            detail["telecaller_comment"] = getattr(record, "telecaller_comment", None)
        if "telecaller_called_at" not in detail or not detail["telecaller_called_at"]:
            called = getattr(record, "telecaller_called_at", None)
            detail["telecaller_called_at"] = called.isoformat() if hasattr(called, "isoformat") else (str(called) if called else None)
        if "telecaller_by_name" not in detail or not detail["telecaller_by_name"]:
            detail["telecaller_by_name"] = (
                (record.telecaller_by.full_name or record.telecaller_by.email)
                if getattr(record, "telecaller_by", None)
                else None
            )

        # Unified QC Status (Merged Audit Decision & QA Status)
        audit_val = detail.get("audit_status")
        if audit_val:
            audit_label_map = {
                "APPROVAL": "Approval",
                "DISAPPROVAL": "Disapproval",
                "COMMERCIAL_PROPERTY": "Commercial Property",
                "RE_AUDITED": "Reaudited",
                "PENDING": "Pending",
            }
            unified_qc = audit_label_map.get(str(audit_val).upper(), str(audit_val).replace("_", " ").title())
            if getattr(record, "qc_edited_at", None):
                unified_qc += " (QC Edited)"
            detail["qc_status"] = unified_qc
            detail["audit_status"] = unified_qc
        elif getattr(record, "qc_edited_at", None):
            detail["qc_status"] = "QC Edited"
            detail["audit_status"] = "QC Edited"
        elif getattr(record, "qc_completed_at", None) or any(
            getattr(record, attr, None) and str(getattr(record, attr, None)).strip()
            for attr in ("qc_vhr_comment", "qc_street_details", "qc_builder_contact", "qc_sqfeet")
        ):
            detail["qc_status"] = "QC Done"
            detail["audit_status"] = "QC Done"
        else:
            detail["qc_status"] = "QC Pending"
            detail["audit_status"] = "QC Pending"

        # Verification Status label
        if detail.get("status") in ("COMPLETED", "Updated Revisited", "UPDATED_REVISITED") or detail.get("revisit_status") == "COMPLETED":
            detail["status"] = "Completed"
        elif detail.get("status") == "NOT_COMPLETED":
            if (
                detail.get("reassigned_to_id") is not None
                or detail.get("reassigned_to_name") is not None
                or detail.get("revisit_status") in ("REASSIGNED", "PENDING_REVISIT")
            ):
                detail["status"] = "Reassigned"
            else:
                detail["status"] = "Pending"
        elif detail.get("status"):
            detail["status"] = str(detail.get("status")).replace("_", " ").title()

        # Format row values
        row_values = [_format_cell_value(f, detail.get(f)) for f in selected_fields]
        sheet.append(row_values)
        row_idx = sheet.max_row
        is_even = (row_idx % 2 == 0)
        row_fill = ROW_FILL_EVEN if is_even else ROW_FILL_ODD

        max_text_len = 0
        for col_idx, (field, cell) in enumerate(zip(selected_fields, sheet[row_idx]), start=1):
            cell.font = DATA_FONT_BOLD if field in ("unique_id", "old_unique_id") else DATA_FONT
            cell.fill = row_fill
            cell.border = DATA_BORDER
            if isinstance(cell.value, str):
                cell.data_type = 's'
                max_text_len = max(max_text_len, len(cell.value))
            if field in CENTER_ALIGNED_FIELDS:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

        # Generous dynamic row height for readability
        sheet.row_dimensions[row_idx].height = max(26, min(75, 20 + (max_text_len // 70) * 14))

        # Photos sheet rows
        if include_images and photos_sheet is not None:
            for photo, metadata in zip(record.images, images):
                photos_sheet.append([
                    record.id, record.unique_id, photo.id, photo.photo_order, photo.file_name,
                    photo.source, metadata.get("captured_at"), metadata.get("created_at"),
                    photo.latitude, photo.longitude, photo.location_accuracy, "Image unavailable",
                ])
                p_row = photos_sheet.max_row
                p_fill = ROW_FILL_EVEN if (p_row % 2 == 0) else ROW_FILL_ODD
                for c in photos_sheet[p_row]:
                    c.font = DATA_FONT
                    c.fill = p_fill
                    c.border = DATA_BORDER
                    c.alignment = ALIGN_CENTER
                    if isinstance(c.value, str):
                        c.data_type = 's'
                photos_sheet.row_dimensions[p_row].height = 26
                try:
                    content = _photo_bytes(record, photo)
                    if content:
                        from PIL import Image
                        with Image.open(io.BytesIO(content)) as original:
                            original.thumbnail((480, 300))
                            preview = io.BytesIO()
                            original.convert("RGB").save(preview, format="JPEG", quality=90)
                        preview.seek(0)
                        image = ExcelImage(preview)
                        photos_sheet.add_image(image, f"L{p_row}")
                        photos_sheet.cell(p_row, 12).value = None
                        photos_sheet.row_dimensions[p_row].height = max(40, image.height * 0.75 + 12)
                except Exception:
                    photos_sheet.cell(p_row, 12).value = "Image unavailable"

    # Modern Sheet View Configurations
    # Freeze Record ID, Unique ID, and Old Unique ID on left, and Header Row on top
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"

    if include_images and photos_sheet is not None:
        photos_sheet.freeze_panes = "D2"
        photos_sheet.auto_filter.ref = photos_sheet.dimensions
        photos_sheet.sheet_view.showGridLines = True
        photos_sheet.column_dimensions["L"].width = 70

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
