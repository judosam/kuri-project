import io
import pytest
import openpyxl

from app.models.user import User, UserRole
from app.core.security import hash_password


def create_test_excel(headers, rows, sheet_name="Task Dispatch"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


class TestExcelTaskDispatch:
    def test_download_template_is_valid_xlsx(self, client, admin_token):
        """Template endpoint must return a styled .xlsx workbook with operator_email."""
        resp = client.get(
            "/api/admin/records/template",
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert "task_dispatch_template.xlsx" in resp.headers.get("content-disposition", "")

        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
        ws = wb.active
        assert ws.title == "Task Dispatch"

        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 3  # Header + 2 sample rows

        headers = rows[0]
        assert "operator_email" in headers
        assert "coordinates" in headers
        assert "site_address" in headers
        assert "unique_id" not in headers
        assert "user_id" not in headers  # Strictly removed in favor of email

    def test_export_records_is_valid_xlsx(self, client, admin_token, test_user):
        """Export endpoint must return a styled .xlsx report."""
        resp = client.get(
            "/api/admin/records/export",
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert resp.headers.get("content-disposition", "").endswith('.xlsx"')

        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
        ws = wb.active
        headers = list(ws.iter_rows(values_only=True))[0]
        assert "Operator Email" in headers
        assert "Field Verifier" in headers
        assert "Imported Coordinates" in headers
        assert "Site Address" in headers

    def test_upload_rejects_csv_with_helpful_message(self, client, admin_token):
        """Uploading .csv must be rejected with Google Sheets / Excel guidance."""
        csv_content = b"operator_email,unique_id,coordinates\ntest@example.com,UID-1,12.9,77.5"
        resp = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.csv", io.BytesIO(csv_content), "text/csv")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 400
        assert "CSV format is not supported" in resp.json()["detail"]
        assert "Google Sheets" in resp.json()["detail"]

    def test_upload_rejects_corrupted_file(self, client, admin_token):
        """Uploading corrupt non-xlsx file named .xlsx must return 400."""
        junk_content = b"Not a real zip or excel file at all"
        resp = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.xlsx", io.BytesIO(junk_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 400
        assert "Unable to read Excel file" in resp.json()["detail"]

    def test_upload_rejects_unrelated_sheet(self, client, admin_token):
        """Uploading an Excel sheet with unrelated headers (e.g. invoice) triggers defense."""
        headers = ["Invoice_Number", "Vendor_Name", "Billing_Amount", "Date"]
        rows = [["INV-001", "Acme Corp", 5000, "2025-01-01"]]
        xlsx_bytes = create_test_excel(headers, rows, sheet_name="InvoiceSheet")

        resp = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("invoices.xlsx", io.BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 400
        assert "Unrecognized sheet format" in resp.json()["detail"]

    def test_upload_successful_by_operator_email(self, client, admin_token, test_user):
        """Tasks are created STAGED for audit: audit_status PENDING, user_id
        NULL, intended operator parked in staged_operator_id until an auditor
        approves. Unique ID is auto-generated from coordinates."""
        headers = [
            "operator_email", "visit_order",
            "coordinates", "street_address", "land_use", "survey_no", "subdivision_no",
            "owner_name_en", "site_address", "survey_number", "sub_division", "taluk",
            "village", "district"
        ]
        rows = [
            [
                test_user.email, 1,
                "12.971598, 77.594566", "100 Feet Rd", "Commercial", "45", "A",
                "Anand Sharma", "100 Feet Road, Indiranagar, Bengaluru", "45", "A",
                "Bangalore East", "Indiranagar", "Bengaluru Urban"
            ],
            [
                test_user.email, 2,
                "12.935242, 77.624462", "Hosur Rd", "Residential", "12", "1",
                "Priya Nair", "Hosur Main Road, Koramangala, Bengaluru", "12", "1",
                "Bangalore South", "Koramangala", "Bengaluru Urban"
            ],
        ]
        xlsx_bytes = create_test_excel(headers, rows)

        resp = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.xlsx", io.BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["success"] is True
        assert data["total_dispatched"] == 2
        assert data["failed_count"] == 0
        assert len(data["errors"]) == 0
        assert len(data["dispatched_tasks"]) == 2
        assert data["dispatched_tasks"][0]["operator_email"] == test_user.email
        assert data["dispatched_tasks"][0]["unique_id"] == "CBE-129716-775946"

        # The field-inspection address means "where the operator stood", so a
        # freshly dispatched task has none — the admin's text stays in
        # site_address. Copying it here made the two addresses identical.
        detail = client.get(
            f"/api/admin/records/{data['dispatched_tasks'][0]['record_id']}",
            headers={"Authorization": f"Bearer {admin_token}"},
        ).json()
        assert detail["formatted_address"] is None
        assert detail["site_address"] == "100 Feet Road, Indiranagar, Bengaluru"

        # Audit workflow: dispatched tasks are staged, not assigned.
        assert detail["audit_status"] == "PENDING"
        assert detail["user_id"] is None
        assert detail["staged_operator_id"] == test_user.id
        assert data["dispatched_tasks"][0]["staged"] is True

    def test_upload_reports_row_errors_for_unknown_operator_and_bad_coords(self, client, admin_token, test_user):
        """Reports specific row-level errors for non-existent operators and bad coordinates."""
        headers = ["operator_email", "unique_id", "coordinates", "site_address"]
        rows = [
            # Row 2: valid
            [test_user.email, "TASK-OK", "12.97, 77.59", "Valid Address"],
            # Row 3: non-existent email
            ["nobody_exists_12345@domain.com", "TASK-ERR-1", "12.97, 77.59", "Address 2"],
            # Row 4: invalid coordinates format
            [test_user.email, "TASK-ERR-2", "not_a_coordinate", "Address 3"],
            # Row 5: coordinates out of range
            [test_user.email, "TASK-ERR-3", "195.0, 77.59", "Address 4"],
        ]
        xlsx_bytes = create_test_excel(headers, rows)

        resp = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.xlsx", io.BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["total_dispatched"] == 1  # only first row succeeded
        assert data["failed_count"] == 3
        assert len(data["errors"]) == 3

        # Verify error rows & columns
        err_msgs = {e["row"]: (e["column"], e["message"]) for e in data["errors"]}
        assert 3 in err_msgs
        assert err_msgs[3][0] == "operator_email"
        assert "nobody_exists_12345@domain.com" in err_msgs[3][1]

        assert 4 in err_msgs
        assert err_msgs[4][0] == "coordinates"

        assert 5 in err_msgs
        assert err_msgs[5][0] == "coordinates"
        assert "out of range" in err_msgs[5][1]

    def test_upload_skips_duplicates_with_warning(self, client, admin_token, test_user):
        """Duplicates (intra-sheet or existing in DB) are skipped with clear warning."""
        headers = [
            "operator_email", "visit_order",
            "coordinates", "street_address", "site_address"
        ]
        # Upload first batch
        rows_1 = [
            [test_user.email, 1, "13.082680, 80.270718", "Poonamallee High Rd", "Chennai Central"],
        ]
        resp1 = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.xlsx", io.BytesIO(create_test_excel(headers, rows_1)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp1.status_code == 200
        assert resp1.json()["total_dispatched"] == 1

        # Upload second batch:
        # Row 1: same coordinates as existing DB record
        # Row 2: new coordinates
        # Row 3: intra-sheet duplicate of Row 2
        rows_2 = [
            [test_user.email, 1, "13.082680, 80.270718", "Poonamallee High Rd", "Chennai Central"],
            [test_user.email, 2, "13.041800, 80.250700", "Anna Salai", "T Nagar"],
            [test_user.email, 3, "13.041800, 80.250700", "Anna Salai", "T Nagar Dup"],
        ]
        resp2 = client.post(
            "/api/admin/records/upload-sheet",
            files={"file": ("tasks.xlsx", io.BytesIO(create_test_excel(headers, rows_2)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        assert resp2.status_code == 200
        data2 = resp2.json()
        assert data2["total_dispatched"] == 1
        assert data2["skipped_duplicates_count"] == 2
        assert len(data2["duplicate_warnings"]) == 2
