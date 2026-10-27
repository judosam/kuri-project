import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

from app.models.capture_record import CaptureRecord
from app.models.user import User, UserRole
from tests.conftest import auth_header
from tests.test_audit_gateway import manager, manager_token, auditor, auditor_token, upload_tasks


def test_three_stage_handoff_and_manual_counts(client, db_session, test_user, user_token, admin_token,
                                              auditor, auditor_token):
    test_user.role = UserRole.FIELD_VERIFIER
    db_session.commit()
    staged = upload_tasks(client, admin_token, [[test_user.email, 'WORKFLOW-1', '11.04, 77.03', 'Site Street']])
    assert staged.status_code == 200, staged.text
    rid = staged.json()['dispatched_tasks'][0]['record_id']
    assert client.get('/api/records', headers=auth_header(user_token)).json()['total'] == 0
    routed = client.post('/api/admin/records/assign-auditor', json={'record_ids': [rid, rid], 'auditor_id': auditor.id}, headers=auth_header(admin_token))
    assert routed.json()['reassigned_count'] == 1
    qc = client.patch(f'/api/admin/records/{rid}/qc', json={'qc_vhr_comment': 'Construction visible'}, headers=auth_header(auditor_token))
    assert qc.status_code == 200 and qc.json()['audit_status'] == 'PENDING'
    decision = client.patch(f'/api/admin/records/{rid}/audit', json={'audit_status': 'APPROVAL'}, headers=auth_header(admin_token))
    assert decision.status_code == 200, decision.text
    assert decision.json()['user_id'] is None
    assert decision.json()['qc_completed_at']
    old = db_session.get(CaptureRecord, rid)
    old.assigned_at = datetime.now(timezone.utc) - timedelta(days=10)
    db_session.commit()
    released = client.post('/api/admin/records/assign-verifier', json={'record_ids': [rid, rid], 'verifier_id': test_user.id}, headers=auth_header(admin_token))
    assert released.json()['reassigned_count'] == 1
    details = client.get(f'/api/records/{rid}', headers=auth_header(user_token)).json()
    assert datetime.fromisoformat(details['assigned_at'].replace('Z', '+00:00')) > datetime.now(timezone.utc) - timedelta(minutes=1)
    done = client.post('/api/records', json={'task_id': rid, 'name': 'WORKFLOW-1', 'photos': [], 'construction_happening': 'Yes'}, headers=auth_header(user_token))
    assert done.status_code == 201, done.text
    assert done.json()['is_manually_added'] is False
    assert done.json()['field_captured_at']  # Completion time is recorded even without GPS permission.
    manual = client.post('/api/records', json={'name': 'Manual visit', 'photos': []}, headers=auth_header(user_token))
    assert manual.status_code == 201, manual.text
    assert manual.json()['is_manually_added'] is True
    for prefix, token in [('/api/admin', admin_token), ('/api/records', user_token)]:
        stats = client.get(f'{prefix}/dashboard', headers=auth_header(token)).json()
        assert stats['manually_added_records'] == 1
        assert stats['total_completed_records'] == 2
    again = client.post('/api/admin/records/assign-verifier', json={'record_ids': [rid], 'verifier_id': test_user.id}, headers=auth_header(admin_token))
    assert again.json()['reassigned_count'] == 1
    blocked = client.patch(f'/api/admin/records/{rid}/audit', json={'audit_status': 'DISAPPROVAL'}, headers=auth_header(auditor_token))
    assert blocked.status_code == 403
    assert client.get(f'/api/records/{rid}', headers=auth_header(user_token)).json()['user_id'] == test_user.id


def test_blocked_direct_access_and_bad_task_never_creates_duplicates(client, db_session, test_user, user_token, admin_token):
    blocked = CaptureRecord(user_id=test_user.id, name='Blocked legacy row', audit_status='DISAPPROVAL')
    db_session.add(blocked)
    db_session.commit()
    for path in [f'/api/records/{blocked.id}', f'/api/records/{blocked.id}/download']:
        assert client.get(path, headers=auth_header(user_token)).status_code == 403
    count = db_session.query(CaptureRecord).count()
    for task_id in [blocked.id, 999999]:
        response = client.post('/api/records', json={'task_id': task_id, 'name': 'Do not duplicate', 'photos': []}, headers=auth_header(user_token))
        assert response.status_code == 404
    assert db_session.query(CaptureRecord).count() == count
    stats = client.get('/api/records/dashboard', headers=auth_header(user_token)).json()
    assert stats['total_records'] == 0 and stats['recent_records'] == []


def test_auditor_routing_and_qc_clearing(client, db_session, auditor, auditor_token, test_admin, admin_token):
    other = User(email='another-auditor@example.com', full_name='Another Auditor', password_hash=auditor.password_hash, role=UserRole.AUDITOR)
    db_session.add(other)
    db_session.flush()
    record = CaptureRecord(name='Other auditor task', user_id=None, assigned_auditor_id=other.id)
    db_session.add(record)
    db_session.commit()
    endpoint = f'/api/admin/records/{record.id}/audit'
    assert client.patch(endpoint, json={'audit_status': 'APPROVAL'}, headers=auth_header(auditor_token)).status_code == 403
    record.assigned_auditor_id = auditor.id
    db_session.commit()
    saved = client.patch(f'/api/admin/records/{record.id}/qc', json={'qc_vhr_comment': 'Confirmed construction', 'qc_builder_contact': 'Builder phone'}, headers=auth_header(auditor_token))
    assert saved.status_code == 200 and saved.json()['qc_completed_at']
    cleared = client.patch(endpoint, json={'audit_status': 'APPROVAL', 'qc_vhr_comment': '', 'qc_street_details': None, 'qc_builder_contact': ''}, headers=auth_header(admin_token))
    assert cleared.status_code == 200
    assert cleared.json()['qc_completed_at'] is None
    assert cleared.json()['qc_builder_contact'] is None
    assert client.delete(f'/api/admin/users/{auditor.id}', headers=auth_header(admin_token)).status_code == 409


def test_admin_template_export_and_no_history_bypass(client, db_session, admin_token, manager_token):
    assert client.get('/api/admin/records/template', headers=auth_header(admin_token)).status_code == 200
    for state in ['APPROVAL', 'DISAPPROVAL']:
        db_session.add(CaptureRecord(user_id=None, name=state, audit_status=state, qc_vhr_comment='Desk notes'))
    db_session.commit()
    exported = client.get('/api/admin/records/export', params={'audit_status': 'DISAPPROVAL', 'qc_status': 'DONE'}, headers=auth_header(admin_token))
    assert exported.status_code == 200, exported.text
    sheet = openpyxl.load_workbook(io.BytesIO(exported.content)).worksheets[0]
    assert sheet.max_row == 2
    fields = [cell.value for cell in sheet[1]]
    assert ('QC Comments (VHR Imagery)' in fields or 'Qc Vhr Comment' in fields)
    decision_col = 'QC Status' if 'QC Status' in fields else ('Audit Decision' if 'Audit Decision' in fields else 'Audit Status')
    assert sheet.cell(2, fields.index(decision_col) + 1).value in ('Disapproval', 'DISAPPROVAL')
    # Retiring the importer role did not hand sheet import to the other office roles.
    bypass = client.post('/api/admin/records/upload-sheet', data={'mode': 'history'}, files={'file': ('old.xlsx', b'not read')}, headers=auth_header(manager_token))
    assert bypass.status_code == 403


def test_legacy_manual_classification_excludes_import_and_task_metadata(client, db_session, test_user, user_token):
    db_session.add_all([
        CaptureRecord(user_id=test_user.id, name='Legacy manual', source='captured', coordinates='11.04, 77.03', status='COMPLETED'),
        CaptureRecord(user_id=test_user.id, name='Imported history', source='imported', status='COMPLETED'),
        CaptureRecord(user_id=test_user.id, name='Legacy dispatched', coordinates='11.04, 77.03', site_address='Dispatched Site', status='COMPLETED'),
    ])
    db_session.commit()
    stats = client.get('/api/records/dashboard', headers=auth_header(user_token)).json()
    assert stats['manually_added_records'] == 1 and stats['total_completed_records'] == 3
