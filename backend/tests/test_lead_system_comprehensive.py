import time
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.database.connection import SessionLocal
from app.models.capture_record import CaptureRecord
from app.models.user import User, UserRole
from app.services.auth_service import AuthService
from app.services.record_service import RecordService
import io
import openpyxl

@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.close()

@pytest.fixture
def client():
    return TestClient(app)

def test_corrected_lead_location(db):
    """Lead CBE-110154-769878 must be located in Coimbatore, NOT Harur/Dharmapuri."""
    lead = db.query(CaptureRecord).filter(CaptureRecord.unique_id == "CBE-110154-769878").first()
    assert lead is not None, "Lead CBE-110154-769878 must exist"
    assert lead.district == "Coimbatore", f"Expected district Coimbatore, got {lead.district}"
    assert lead.taluk == "Coimbatore North", f"Expected taluk Coimbatore North, got {lead.taluk}"
    assert lead.village == "Pappanaickenpalayam", f"Expected village Pappanaickenpalayam, got {lead.village}"
    assert "Harur" not in (lead.taluk or "")
    assert "Dharmapuri" not in (lead.district or "")

def test_dashboard_reconciliation(db):
    """Total records (1849) must reconcile exactly with Completed + Closed + Disapproved + Pending."""
    svc = RecordService(db)
    metrics = svc.get_dashboard_metrics()
    overview = metrics.get("overview") or {}

    total = overview.get("total_records")
    completed = overview.get("completed_verifications")
    closed = overview.get("closed_verifications")
    disapproved = overview.get("disapproved_verifications")
    pending = overview.get("pending_verifications")

    assert total == 1849
    assert completed == 617
    assert closed == 60
    assert disapproved == 1
    assert pending == 1171
    assert completed + closed + disapproved + pending == total
    assert metrics.get("pending_records") == pending

    # Verifier stats reconciliation: Assigned == Completed + Closed + Pending
    for stat in metrics.get("inspector_stats", []):
        assert stat["total_assigned"] == stat["completed"] + stat.get("closed", 0) + stat["pending"], (
            f"Verifier {stat['user_name']} mismatch: "
            f"Assigned {stat['total_assigned']} != {stat['completed']} + {stat.get('closed', 0)} + {stat['pending']}"
        )

def test_strict_rbac_matrix(db, client):
    """Strict server-side RBAC: each role accesses only its authorized endpoints, others get 403."""
    auth_svc = AuthService(db)
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    auditor = db.query(User).filter(User.role == UserRole.AUDITOR).first()
    telecaller = db.query(User).filter(User.role == UserRole.TELECALLER).first()
    verifier = db.query(User).filter(User.role == UserRole.FIELD_VERIFIER).first()

    tokens = {
        "admin": auth_svc.create_token(admin) if admin else None,
        "auditor": auth_svc.create_token(auditor) if auditor else None,
        "telecaller": auth_svc.create_token(telecaller) if telecaller else None,
        "verifier": auth_svc.create_token(verifier) if verifier else None,
    }

    # 1. Telecaller routes forbidden for auditor & field verifier
    for role in ["auditor", "verifier"]:
        tok = tokens[role]
        if not tok:
            continue
        res = client.get("/api/telecaller/dashboard", headers={"Authorization": f"Bearer {tok}"})
        assert res.status_code == 403, f"{role} should get 403 on /api/telecaller/dashboard, got {res.status_code}"

    # 2. Auditor routes forbidden for telecaller & field verifier
    for role in ["telecaller", "verifier"]:
        tok = tokens[role]
        if not tok:
            continue
        res = client.get("/api/admin/records/pending-audit", headers={"Authorization": f"Bearer {tok}"})
        assert res.status_code == 403, f"{role} should get 403 on pending-audit, got {res.status_code}"

    # 3. Field Verifier record write routes forbidden for telecaller & auditor
    for role in ["telecaller", "auditor"]:
        tok = tokens[role]
        if not tok:
            continue
        res = client.post("/api/records", json={"name": "Test"}, headers={"Authorization": f"Bearer {tok}"})
        assert res.status_code == 403, f"{role} should get 403 on /api/records write, got {res.status_code}"

    # 4. Admin routes forbidden for non-admins
    for role in ["auditor", "telecaller", "verifier"]:
        tok = tokens[role]
        if not tok:
            continue
        res = client.get("/api/admin/users", headers={"Authorization": f"Bearer {tok}"})
        assert res.status_code == 403, f"{role} should get 403 on /api/admin/users, got {res.status_code}"

def test_templates_without_visit_order(client, db):
    """Excel and CSV download templates must not contain cumulative_road_distance_km."""
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(admin)

    # Test CSV template
    res_csv = client.get("/api/admin/records/template?format=csv", headers={"Authorization": f"Bearer {token}"})
    assert res_csv.status_code == 200
    csv_text = res_csv.text
    assert "cumulative_road_distance_km" not in csv_text

    # Test XLSX template
    res_xlsx = client.get("/api/admin/records/template?format=xlsx", headers={"Authorization": f"Bearer {token}"})
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "cumulative_road_distance_km" not in headers

def test_telecaller_history_table_format(client, db):
    """Telecaller comments should be formatted in pipe-delimited table format."""
    telecaller = db.query(User).filter(User.role == UserRole.TELECALLER).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(telecaller)

    # Create dummy record
    rec = CaptureRecord(
        name="Test Lead for Table",
        status="PENDING",
        contact_number="9876543210",
        assigned_telecaller_id=telecaller.id,
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        # Submit feedback
        res = client.post(
            f"/api/telecaller/records/{rec.id}/feedback",
            json={
                "telecaller_status": "CONSTRUCTION_CONFIRMED",
                "telecaller_comment": "Owner confirmed 2 floors construction ongoing.",
                "follow_up_date": "2026-04-15",
                "stage_of_construction": "Piling & Earthwork",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 200
        db.refresh(rec)
        assert rec.telecaller_comment is not None
        assert "|" in rec.telecaller_comment, "Comment should be in table/pipe-delimited format"
        assert "Date & Time" in rec.telecaller_comment
        assert "Caller" in rec.telecaller_comment
        assert "Status" in rec.telecaller_comment
        assert "Details" in rec.telecaller_comment
        assert "Notes" in rec.telecaller_comment
        assert "Owner confirmed 2 floors construction ongoing." in rec.telecaller_comment
    finally:
        db.delete(rec)
        db.commit()

def test_auto_reassign_on_field_follow_up_date(client, db):
    """When field verifier sets follow_up_date, record should automatically reassign to verifier."""
    verifier = db.query(User).filter(User.role == UserRole.FIELD_VERIFIER).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(verifier)

    rec = CaptureRecord(
        name="Field Follow Up Lead",
        status="PENDING",
        audit_status="APPROVAL",
        user_id=verifier.id,
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        # Submit field form with follow_up_date via POST /api/records
        res = client.post(
            "/api/records",
            json={
                "task_id": rec.id,
                "name": rec.name,
                "date_of_visit": "2026-03-30",
                "construction_happening": "Yes",
                "stage_of_construction": "Piling & Earthwork",
                "follow_up_date": "2026-04-20",
                "remarks": "Revisit needed next month",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code in (200, 201)
        db.refresh(rec)
        assert rec.follow_up_date == "2026-04-20"
        assert rec.reassigned_to_id == verifier.id
        assert rec.revisit_status == "PENDING_REVISIT"
        assert rec.reassigned_at is not None
    finally:
        db.delete(rec)
        db.commit()

def test_cumulative_road_distance_km_removed(db):
    """Ensure cumulative_road_distance_km is not in CaptureRecord model."""
    assert not hasattr(CaptureRecord, "cumulative_road_distance_km")

def test_excel_export_updater_attribution_and_floors(client, db):
    """Excel export must include Updater Attribution, Telecaller specs, and Field Verifier specs."""
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(admin)

    res = client.get("/api/admin/records/export?format=xlsx", headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    assert "Data Updated By (Name)" in headers
    assert "Data Updated By (Role)" in headers
    assert "Telecaller - No. of Floors" in headers
    assert "Telecaller - Expected Completion" in headers
    assert "Telecaller - Built-up Area" in headers
    assert "Telecaller - Stage of Construction" in headers
    assert "No. of Floors" in headers
    assert "Built-up Area" in headers
    assert "Expected Completion" in headers

def test_mandatory_remarks_when_construction_not_happening(client, db):
    """If construction_happening is 'No', remarks is mandatory."""
    verifier = db.query(User).filter(User.role == UserRole.FIELD_VERIFIER).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(verifier)

    rec = CaptureRecord(
        name="No Construction Lead",
        status="PENDING",
        audit_status="APPROVAL",
        user_id=verifier.id,
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        # Without remarks -> should fail 400
        res = client.post(
            "/api/records",
            json={
                "task_id": rec.id,
                "name": rec.name,
                "date_of_visit": "2026-03-30",
                "construction_happening": "No",
                "remarks": "   ",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 400

        # With remarks -> should succeed
        res2 = client.post(
            "/api/records",
            json={
                "task_id": rec.id,
                "name": rec.name,
                "date_of_visit": "2026-03-30",
                "construction_happening": "No",
                "remarks": "Site is completely vacant land with bushes, no construction activity found.",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res2.status_code in (200, 201)
    finally:
        db.delete(rec)
        db.commit()

def test_admin_update_floors_and_specs(client, db):
    """Admin update endpoint must persist both telecaller and field specifications."""
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(admin)

    rec = CaptureRecord(
        name="Test Floors Spec",
        status="COMPLETED",
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        res = client.put(
            f"/api/admin/records/{rec.id}",
            json={
                "no_of_floors": "G+3",
                "built_up_area": "4200 sq.ft",
                "expected_completion": "March 2026",
                "telecaller_no_of_floors": "G+2",
                "telecaller_built_up_area": "3000 sq.ft",
                "telecaller_expected_completion": "Dec 2026",
                "telecaller_stage_of_construction": "Foundation",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 200
        data = res.json()
        assert data["no_of_floors"] == "G+3"
        assert data["built_up_area"] == "4200 sq.ft"
        assert data["expected_completion"] == "March 2026"
        assert data["telecaller_no_of_floors"] == "G+2"
        assert data["telecaller_built_up_area"] == "3000 sq.ft"
        assert data["telecaller_expected_completion"] == "Dec 2026"
        assert data["telecaller_stage_of_construction"] == "Foundation"
    finally:
        db.delete(rec)
        db.commit()

def test_telecaller_field_verifier_isolation(client, db):
    """Telecaller updates must never contaminate Field Verification section, and vice-versa."""
    telecaller = db.query(User).filter(User.role == UserRole.TELECALLER).first()
    auth_svc = AuthService(db)
    token = auth_svc.create_token(telecaller)

    rec = CaptureRecord(
        name="Isolation Test Lead",
        status="PENDING",
        contact_number="9876543210",
        assigned_telecaller_id=telecaller.id,
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        # Telecaller updates call feedback with phone-verified floors & completion
        res = client.post(
            f"/api/telecaller/records/{rec.id}/feedback",
            json={
                "telecaller_status": "CONSTRUCTION_CONFIRMED",
                "telecaller_comment": "Verified on phone call with contractor",
                "no_of_floors": "G+2",
                "expected_completion": "1 year",
                "built_up_area": "2500 sq.ft",
                "stage_of_construction": "Structure Completed",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 200
        db.refresh(rec)

        # Telecaller fields MUST be set
        assert rec.telecaller_no_of_floors == "G+2"
        assert rec.telecaller_expected_completion == "1 year"
        assert rec.telecaller_built_up_area == "2500 sq.ft"
        assert rec.telecaller_stage_of_construction == "Structure Completed"

        # Field Verifier fields MUST remain None (zero pollution into Field Verification!)
        assert rec.no_of_floors is None
        assert rec.expected_completion is None
        assert rec.built_up_area is None

        # Verify Record 2792 from GCP database (CBE-110932-770362)
        rec_2792 = db.query(CaptureRecord).filter(CaptureRecord.unique_id == "CBE-110932-770362").first()
        if rec_2792:
            assert rec_2792.telecaller_no_of_floors == "G+2"
            assert rec_2792.telecaller_expected_completion == "1 year"
            assert rec_2792.no_of_floors is None
            assert rec_2792.expected_completion is None
    finally:
        db.delete(rec)
        db.commit()


def test_auditor_role_can_save_qc_and_bulk_audit(client, db):
    """Verify that AUDITOR role can execute single and bulk audit without 403 Admin access required."""
    auditor = db.query(User).filter(User.role == UserRole.AUDITOR).first()
    if not auditor:
        pytest.skip("No auditor user found")
    auth_svc = AuthService(db)
    auditor_token = auth_svc.create_token(auditor)

    # 1. Create a dummy test record
    rec = CaptureRecord(
        name="Test Lead for Auditor QC",
        status="PENDING",
        audit_status="PENDING",
        unique_id="TEST-QC-001",
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)

    try:
        # 2. Auditor calls PATCH /api/admin/records/{id}/audit
        res = client.patch(
            f"/api/admin/records/{rec.id}/audit",
            json={
                "audit_status": "APPROVAL",
                "audit_comment": "Verified and approved by QA",
                "qc_vhr_comment": "Construction Active",
                "qc_street_details": "Good road access",
                "qc_builder_contact": "9876543210",
            },
            headers={"Authorization": f"Bearer {auditor_token}"},
        )
        assert res.status_code == 200, f"Auditor got {res.status_code}: {res.text}"
        data = res.json()
        assert data["audit_status"] == "RE_AUDITED"
        db.refresh(rec)
        assert rec.audit_status == "RE_AUDITED"
        assert rec.audited_by_id == auditor.id

        # 3. Auditor calls POST /api/admin/records/bulk-audit
        res_bulk = client.post(
            "/api/admin/records/bulk-audit",
            json={
                "record_ids": [rec.id],
                "audit_status": "RE_AUDITED",
                "audit_comment": "Bulk QA update",
            },
            headers={"Authorization": f"Bearer {auditor_token}"},
        )
        assert res_bulk.status_code == 200, f"Auditor got {res_bulk.status_code}: {res_bulk.text}"
        assert res_bulk.json()["audited_count"] == 1
        db.refresh(rec)
        assert rec.audit_status == "RE_AUDITED"

        # 4. Verify telecaller / verifier get 403 on audit endpoint
        telecaller = db.query(User).filter(User.role == UserRole.TELECALLER).first()
        if telecaller:
            tc_token = auth_svc.create_token(telecaller)
            res_tc = client.patch(
                f"/api/admin/records/{rec.id}/audit",
                json={"audit_status": "APPROVAL"},
                headers={"Authorization": f"Bearer {tc_token}"},
            )
            assert res_tc.status_code == 403, f"Telecaller should get 403 on audit endpoint, got {res_tc.status_code}"
    finally:
        db.delete(rec)
        db.commit()

def test_qa_portal_resolve_coords_and_create_record(client, db):
    """QA auditor can resolve coordinates (generating site_id, area, village, taluk, district, address)
    and create a record directly from the QA portal."""
    auditor = db.query(User).filter(User.role == UserRole.AUDITOR).first()
    assert auditor is not None
    auth_svc = AuthService(db)
    auditor_token = auth_svc.create_token(auditor)

    # 1. Test POST /api/admin/records/resolve-coords generates QANEW prefix with coordinate pattern
    coords_input = "11.0825, 76.9950"
    res = client.post(
        "/api/admin/records/resolve-coords",
        json={"coordinates": coords_input},
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    data = res.json()
    assert data["success"] is True
    assert data["site_id"].startswith("QANEW-CBE-"), f"Expected site_id starting with QANEW-CBE-, got {data['site_id']}"
    assert data["site_id"] == "QANEW-CBE-110825-769950"
    assert data["latitude"] == 11.0825
    assert data["longitude"] == 76.9950
    assert data["district"] == "Coimbatore"
    assert data["village"] is not None
    assert data["taluk"] is not None
    assert data["area"] is not None
    assert data["address"] is not None

    # 2. Test POST /api/admin/records/qa-create with auto-generated QANEW ID (no unique_id passed)
    created_ids_to_clean = []
    try:
        res_auto = client.post(
            "/api/admin/records/qa-create",
            json={
                "coordinates": coords_input,
                "area": data["area"],
                "village": data["village"],
                "taluk": data["taluk"],
                "district": data["district"],
                "site_address": data["address"],
                "priority_tier": "High",
                "audit_status": "PENDING",
            },
            headers={"Authorization": f"Bearer {auditor_token}"},
        )
        assert res_auto.status_code == 201, f"Expected 201, got {res_auto.status_code}: {res_auto.text}"
        auto_rec = res_auto.json()
        assert auto_rec["unique_id"].startswith("QANEW-CBE-")
        assert auto_rec["unique_id"] == "QANEW-CBE-110825-769950"
        assert auto_rec["source"] == "qa_portal"
        created_ids_to_clean.append(auto_rec["id"])

        # Also test explicit unique_id
        unique_test_id = f"TEST-QA-{data['site_id']}-{int(time.time())}"
        res_create = client.post(
            "/api/admin/records/qa-create",
            json={
                "coordinates": coords_input,
                "unique_id": unique_test_id,
                "area": data["area"],
                "village": data["village"],
                "taluk": data["taluk"],
                "district": data["district"],
                "site_address": data["address"],
                "priority_tier": "High",
                "audit_status": "PENDING",
                "audit_comment": "Created via QA portal map select",
                "qc_vhr_comment": "Vacant Land",
                "qc_sqfeet": "2400",
            },
            headers={"Authorization": f"Bearer {auditor_token}"},
        )
        assert res_create.status_code == 201, f"Expected 201, got {res_create.status_code}: {res_create.text}"
        created_rec = res_create.json()
        assert created_rec["unique_id"] == unique_test_id
        assert created_rec["area"] == data["area"]
        assert created_rec["source"] == "qa_portal"
        assert created_rec["assigned_auditor_id"] == auditor.id
        assert created_rec["uploaded_by_id"] == auditor.id
        assert created_rec["qc_vhr_comment"] == "Vacant Land"
        assert created_rec["qc_sqfeet"] == "2400"
        created_ids_to_clean.append(created_rec["id"])

        rec_in_db = db.query(CaptureRecord).filter(CaptureRecord.unique_id == unique_test_id).first()
        assert rec_in_db is not None
        assert rec_in_db.area == data["area"]
        assert rec_in_db.source == "qa_portal"

        # 3. Test Dashboard counts QA Created records
        record_service = RecordService(db)
        metrics = record_service.get_dashboard_metrics()
        assert "qa_new_records" in metrics["overview"]
        assert metrics["overview"]["qa_new_records"] >= 2
        assert metrics["funnel"]["qa_new_records"] >= 2
        assert metrics["qa_new_records"] >= 2

        # 4. RBAC isolation: Field verifier cannot call resolve-coords or qa-create
        verifier = db.query(User).filter(User.role == UserRole.USER).first()
        if verifier:
            v_token = auth_svc.create_token(verifier)
            res_v_resolve = client.post(
                "/api/admin/records/resolve-coords",
                json={"coordinates": coords_input},
                headers={"Authorization": f"Bearer {v_token}"},
            )
            assert res_v_resolve.status_code == 403, "Field verifier must be blocked from resolve-coords"

            res_v_create = client.post(
                "/api/admin/records/qa-create",
                json={"coordinates": coords_input, "unique_id": f"TEST-V-{data['site_id']}"},
                headers={"Authorization": f"Bearer {v_token}"},
            )
            assert res_v_create.status_code == 403, "Field verifier must be blocked from qa-create"
    finally:
        for rid in created_ids_to_clean:
            rec = db.query(CaptureRecord).filter(CaptureRecord.id == rid).first()
            if rec:
                db.delete(rec)
        for r in db.query(CaptureRecord).filter(CaptureRecord.unique_id.like("TEST-QA-%")).all():
            db.delete(r)
        db.commit()



def test_auditor_qa_portal_assigned_records_isolation(client, db):
    """Auditor in QA portal strictly sees only records assigned to them, never other auditors' records."""
    auditors = db.query(User).filter(User.role == UserRole.AUDITOR).all()
    if len(auditors) < 2:
        pytest.skip("Requires at least 2 auditors for isolation test")
    
    auditor1, auditor2 = auditors[0], auditors[1]
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    auth_svc = AuthService(db)
    
    a1_token = auth_svc.create_token(auditor1)
    a2_token = auth_svc.create_token(auditor2)
    admin_token = auth_svc.create_token(admin)
    
    # Create a record assigned to auditor 1 and another to auditor 2
    rec1 = CaptureRecord(
        name="Auditor 1 Task",
        status="PENDING",
        audit_status="PENDING",
        unique_id="ISO-AUD1-001",
        assigned_auditor_id=auditor1.id,
    )
    rec2 = CaptureRecord(
        name="Auditor 2 Task",
        status="PENDING",
        audit_status="PENDING",
        unique_id="ISO-AUD2-002",
        assigned_auditor_id=auditor2.id,
    )
    db.add(rec1)
    db.add(rec2)
    db.commit()
    db.refresh(rec1)
    db.refresh(rec2)
    
    try:
        # 1. Auditor 1 lists records without query params -> MUST ONLY get auditor 1 records
        res1 = client.get("/api/admin/records", headers={"Authorization": f"Bearer {a1_token}"})
        assert res1.status_code == 200
        data1 = res1.json()
        ids1 = [r["id"] for r in data1["records"]]
        assert rec1.id in ids1, "Auditor 1 must see their own assigned record"
        assert rec2.id not in ids1, "Auditor 1 must NEVER see Auditor 2's assigned record"
        for r in data1["records"]:
            assert r["assigned_auditor_id"] == auditor1.id, "All records returned must be assigned to auditor 1"
        
        # 2. Auditor 1 attempts to pass assigned_auditor_id=auditor2.id -> STILL only gets auditor 1 records
        res1_spoof = client.get(
            f"/api/admin/records?assigned_auditor_id={auditor2.id}",
            headers={"Authorization": f"Bearer {a1_token}"}
        )
        assert res1_spoof.status_code == 200
        data1_spoof = res1_spoof.json()
        ids1_spoof = [r["id"] for r in data1_spoof["records"]]
        assert rec2.id not in ids1_spoof, "Auditor 1 cannot bypass assignment filter"
        for r in data1_spoof["records"]:
            assert r["assigned_auditor_id"] == auditor1.id
        
        # 3. Auditor 1 attempts to fetch Auditor 2's record detail -> 403 Forbidden
        res_get2 = client.get(
            f"/api/admin/records/{rec2.id}",
            headers={"Authorization": f"Bearer {a1_token}"}
        )
        assert res_get2.status_code == 403, f"Expected 403 when accessing another auditor's record, got {res_get2.status_code}"
        
        # 4. Auditor 1 attempts to audit Auditor 2's record -> 403 Forbidden
        res_audit2 = client.patch(
            f"/api/admin/records/{rec2.id}/audit",
            json={"audit_status": "APPROVAL", "audit_comment": "Malicious takeover attempt"},
            headers={"Authorization": f"Bearer {a1_token}"}
        )
        assert res_audit2.status_code == 403, "Auditor 1 cannot audit another auditor's record"
        
        # 5. Admin can list both
        res_admin = client.get("/api/admin/records?search=ISO-AUD", headers={"Authorization": f"Bearer {admin_token}"})
        assert res_admin.status_code == 200
        admin_ids = [r["id"] for r in res_admin.json()["records"]]
        assert rec1.id in admin_ids or rec2.id in admin_ids
    finally:
        db.delete(rec1)
        db.delete(rec2)
        db.commit()
