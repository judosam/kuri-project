import io
import zipfile
from datetime import datetime, timezone

import openpyxl
import pytest
from PIL import Image

from app.models.capture_record import CaptureRecord
from app.models.captured_image import CapturedImage
from app.services.local_storage_service import LocalStorageService


@pytest.mark.parametrize('source', ['captured', 'imported'])
def test_admin_full_record_excel(client, admin_token, user_token, test_user, db_session, source):
    record = CaptureRecord(user_id=test_user.id, name='Test inspection', source=source,
                           unique_id='SITE-001', street_address='23 Test Street',
                           survey_no='003', relationship_with_building='Tenant',
                           remarks='=1+1', contact_number='0012345678', status='COMPLETED')
    db_session.add(record)
    db_session.commit()
    image_bytes = io.BytesIO()
    Image.new('RGB', (600, 400), '#447788').save(image_bytes, format='JPEG')
    key = LocalStorageService().save_file(image_bytes.getvalue(), 'excel-test.jpg', f'{test_user.id}/{record.id}')
    photo = CapturedImage(record_id=record.id, google_drive_file_id=key, file_name='excel-test.jpg',
                          photo_order=1, source='uploaded', latitude=0, longitude=0,
                          location_accuracy=0, captured_at=datetime.now(timezone.utc))
    db_session.add(photo)
    db_session.commit()
    url = f'/api/admin/records/{record.id}/export'
    assert client.get(url).status_code == 401
    assert client.get(url, headers={'Authorization': f'Bearer {user_token}'}).status_code == 403
    response = client.get(url, headers={'Authorization': f'Bearer {admin_token}'})
    assert response.status_code == 200
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert response.headers['content-disposition'].endswith(f'Record_{record.id}_full.xlsx"')
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ['Record Details', 'Photos']
    details = {row[0].value: row[1] for row in workbook['Record Details'].iter_rows(min_row=2)}
    assert details['Street Address'].value == '23 Test Street'
    assert details['Survey No'].value == '003'
    assert details['Relationship With Building'].value == 'Tenant'
    assert details['Contact Number'].value == '0012345678'
    assert details['Remarks'].value == '=1+1' and details['Remarks'].data_type == 's'
    assert details['Record Source'].value == source
    assert details['Photo Count'].value == 1
    assert workbook['Photos']['C2'].value == 'uploaded'
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == 1

    report = client.get('/api/admin/records/export', headers={'Authorization': f'Bearer {admin_token}'})
    assert report.status_code == 200
    bulk = openpyxl.load_workbook(io.BytesIO(report.content))
    assert bulk.sheetnames == ['Inspections Report', 'Photos']
    columns = [cell.value for cell in bulk['Inspections Report'][1]]
    values = dict(zip(columns, [cell.value for cell in bulk['Inspections Report'][2]]))
    assert values['Record Name'] == 'Test inspection'
    assert values['Record Source'] == source
    remarks_key = 'Inspection Remarks' if 'Inspection Remarks' in values else 'Remarks'
    assert values[remarks_key] == '=1+1'
    assert values['Photo Count'] == 1
    assert values['Captured At (UTC)'].endswith('Z')
    assert 'Field Inspection Address' in values and 'Admin Resolved Address' in values
    assert bulk['Inspections Report'].cell(2, columns.index(remarks_key) + 1).data_type == 's'
    assert bulk['Photos']['A2'].value == record.id
    assert bulk['Photos']['F2'].value == 'uploaded'
    with zipfile.ZipFile(io.BytesIO(report.content)) as archive:
        assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == 1


def test_excel_without_photos_and_missing_record(client, admin_token, test_user, db_session):
    record = CaptureRecord(user_id=test_user.id, name='No photo inspection')
    db_session.add(record)
    db_session.commit()
    headers = {'Authorization': f'Bearer {admin_token}'}
    response = client.get(f'/api/admin/records/{record.id}/export', headers=headers)
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook['Photos']['A2'].value == 'No photos attached to this record.'
    assert client.get('/api/admin/records/999999/export', headers=headers).status_code == 404


def test_full_report_crosses_pages_and_respects_filters(client, admin_token, user_token, test_user, db_session, monkeypatch):
    import app.services.record_excel_service as service
    monkeypatch.setattr(service, 'REPORT_BATCH_SIZE', 2)
    for index in range(5):
        db_session.add(CaptureRecord(user_id=test_user.id, name=f'Site {index}',
                                    status='COMPLETED' if index < 3 else 'NOT_COMPLETED',
                                    created_at=datetime(2026, 9, 20 + index, tzinfo=timezone.utc)))
    db_session.commit()
    headers = {'Authorization': f'Bearer {admin_token}'}
    url = '/api/admin/records/export'
    assert client.get(url).status_code == 401
    assert client.get(url, headers={'Authorization': f'Bearer {user_token}'}).status_code == 403
    all_rows = client.get(url, headers=headers)
    assert all_rows.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(all_rows.content))
    assert workbook['Inspections Report'].max_row == 6
    params = {'status': 'COMPLETED', 'user_id': test_user.id, 'search': 'Site',
              'date_from': '2026-09-21T00:00:00Z', 'date_to': '2026-09-22T23:59:59Z'}
    selected = client.get(url, headers=headers, params=params)
    assert selected.status_code == 200
    report = openpyxl.load_workbook(io.BytesIO(selected.content))['Inspections Report']
    columns = [cell.value for cell in report[1]]
    name_index = columns.index('Record Name')
    assert {row[name_index].value for row in report.iter_rows(min_row=2)} == {'Site 1', 'Site 2'}
    empty = client.get(url, headers=headers, params={'search':'does-not-exist'})
    assert openpyxl.load_workbook(io.BytesIO(empty.content))['Inspections Report'].max_row == 1
